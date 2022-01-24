'''
Date: 2022-01-24 20:38:32
Author: ywyz
LastModifiedBy: ywyz
Github: https://github.com/ywyz
LastEditors: ywyz
LastEditTime: 2022-01-24 21:15:03
'''
'''
导出幼儿信息排查表中数据
'''
import docx
from docx import Document

path = '幼儿信息排查表20210506副本.docx'
docus = Document(path)
tables = docus.tables
table = tables[0]
print("读取表格成功")
for i in range(0, len(table.rows)):  # 遍历行
    for n in range(0, 13):
        result = table.cell(i, n).text
        print(result, end=" ")
    print()
